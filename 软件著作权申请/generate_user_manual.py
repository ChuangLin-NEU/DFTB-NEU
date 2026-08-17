# -*- coding: utf-8 -*-
"""Generate soft-copyright user manual for DFTB-NEU."""
from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"C:\Users\Focalors\Desktop\DFT-NEU")
OUT_DIR = ROOT / "软件著作权申请"
PPT_MEDIA = ROOT / "_tmp_inspect" / "ppt_media"
UI_IMAGES = ROOT / "_tmp_inspect" / "ui_images"

SOFT_FULL = "DFTB-NEU密度泛函紧束缚计算工作台软件"
SOFT_VER = "V1.0"
HEADER = f"{SOFT_FULL} {SOFT_VER}"


def set_run_font(run, size=12, bold=False, name="宋体", color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, *, size=12, bold=False, align="left", space_after=6, first_line=False, name="宋体"):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line:
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, name=name)
    return p


def add_heading_cn(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 12}
    p = add_para(doc, text, size=sizes.get(level, 12), bold=True, space_after=10, name="黑体")
    return p


def add_body(doc, text):
    return add_para(doc, text, size=12, align="justify", first_line=True, space_after=6)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    pf = p.paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.left_indent = Cm(0.74 + 0.5 * level)
    run = p.add_run(text)
    set_run_font(run, size=12, name="宋体")
    return p


def set_cell_text(cell, text, *, bold=False, size=10.5, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, name="宋体")


def shade_cell(cell, fill="D9E2F3"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, h, bold=True, center=True)
        shade_cell(cell)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            set_cell_text(table.rows[r_i + 1].cells[c_i], val, size=10.5)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_image(doc, path: Path, width_cm=14.5, caption=None):
    if not path.exists():
        add_para(doc, f"【示意图缺失：{path.name}】", size=10.5, align="center", space_after=6)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    if caption:
        add_para(doc, caption, size=10.5, align="center", space_after=10, name="楷体")


def add_page_break(doc):
    doc.add_page_break()


def setup_header_footer(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run(HEADER)
    set_run_font(run, size=9, name="宋体", color=RGBColor(0x55, 0x55, 0x55))

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # PAGE field
    run1 = fp.add_run("— ")
    set_run_font(run1, size=9, name="宋体")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r2 = fp.add_run()
    r2._r.append(fld_begin)
    r2._r.append(instr)
    r2._r.append(fld_sep)
    r2._r.append(fld_text)
    r2._r.append(fld_end)
    set_run_font(r2, size=9, name="宋体")
    run3 = fp.add_run(" —")
    set_run_font(run3, size=9, name="宋体")


def build():
    doc = Document()
    setup_header_footer(doc)

    # ========== Cover ==========
    for _ in range(3):
        add_para(doc, "", space_after=0)
    add_para(doc, "软件用户手册", size=26, bold=True, align="center", space_after=18, name="黑体")
    add_para(doc, SOFT_FULL, size=18, bold=True, align="center", space_after=8, name="黑体")
    add_para(doc, f"版本：{SOFT_VER}", size=14, align="center", space_after=24, name="宋体")
    icon = UI_IMAGES / "icon-student-512.png"
    if icon.exists():
        add_image(doc, icon, width_cm=3.2)
    add_para(doc, "著作权人：东北大学", size=12, align="center", space_after=4)
    add_para(doc, "开发单位：东北大学材料科学与工程学院", size=12, align="center", space_after=4)
    add_para(doc, "文档类型：软件著作权登记申请附件", size=12, align="center", space_after=4)
    add_para(doc, "编写日期：2026年8月", size=12, align="center", space_after=4)

    add_page_break(doc)

    # ========== TOC ==========
    add_heading_cn(doc, "目录", 1)
    toc = [
        "1 软件概述",
        "  1.1 软件全称与版本",
        "  1.2 软件定位与用途",
        "  1.3 系统组成（三端协同）",
        "  1.4 运行环境",
        "2 安装与启动",
        "  2.1 学生端安装",
        "  2.2 教师端安装",
        "  2.3 启动与登录",
        "3 学生端功能说明",
        "  3.1 界面总览",
        "  3.2 首页",
        "  3.3 环境部署",
        "  3.4 固体物理课程",
        "  3.5 自然语言计算",
        "  3.6 任务与结果",
        "  3.7 设置",
        "4 教师端功能说明",
        "5 典型操作流程",
        "6 常见问题与故障排查",
        "7 注意事项与免责声明",
    ]
    for item in toc:
        add_para(doc, item, size=12, space_after=3)

    add_page_break(doc)

    # ========== 1 ==========
    add_heading_cn(doc, "1 软件概述", 1)
    add_heading_cn(doc, "1.1 软件全称与版本", 2)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["软件全称", SOFT_FULL],
            ["软件简称", "DFTB-NEU"],
            ["版本号", SOFT_VER],
            ["软件分类", "应用软件"],
            ["开发完成日期", "2026年8月5日"],
            ["运行平台", "Windows 10/11 x64（计算依托 WSL2 / Ubuntu）"],
            ["主要开发语言", "Python、JavaScript、HTML/CSS、PowerShell"],
        ],
        col_widths=[4, 12],
    )

    add_heading_cn(doc, "1.2 软件定位与用途", 2)
    add_body(
        doc,
        "DFTB-NEU 是面向科研与教学的 DFTB+（密度泛函紧束缚）本机计算工作台。"
        "软件将高水平原子尺度计算能力封装为 Windows 桌面客户端，通过 WSL2 隔离环境完成本机 DFTB+ 计算，"
        "并提供课程训练、自然语言生成输入文件、任务跟踪与结果可视化等功能。"
        "软件强调“本机优先、教学就绪、科研可用”：计算在用户本机完成，不依赖远程超算即可开展课堂演示与课题计算。"
    )
    add_body(
        doc,
        "软件深度融合人工智能大模型能力：用户可用自然语言描述材料结构与计算意图，系统自动生成 DFTB+ 所需的 HSD 输入，"
        "经用户确认后再投递本机计算。该设计降低了 HSD/POSCAR 对初学者的门槛，同时保留人工确认环节，避免误算。"
    )

    add_heading_cn(doc, "1.3 系统组成（三端协同）", 2)
    add_body(doc, "软件由学生端、教师端与课堂中心三部分协同构成：")
    add_table(
        doc,
        ["组成部分", "主要职责"],
        [
            ["学生端（工作台）", "登录、本机部署、课程训练、自然语言计算、任务与观测量、设置"],
            ["教师端", "发布/吊销本堂密码、监管在线学生与会话、结束本堂；无计算职责"],
            ["课堂中心", "鉴权与会话管理，LLM / Materials Project 等代理服务"],
        ],
        col_widths=[4.5, 11.5],
    )
    add_image(
        doc,
        PPT_MEDIA / "image2.png",
        width_cm=14.5,
        caption="图 1-1  DFTB-NEU 三端协同架构示意",
    )

    add_heading_cn(doc, "1.4 运行环境", 2)
    add_para(doc, "（1）硬件环境", bold=True, size=12, space_after=4)
    add_table(
        doc,
        ["项目", "开发环境建议", "运行环境最低要求"],
        [
            ["CPU", "Intel/AMD x64", "x64 处理器"],
            ["内存", "≥16 GB", "≥8 GB（推荐 16 GB）"],
            ["磁盘", "≥50 GB 可用空间", "≥40 GB 可用空间（固态硬盘更佳）"],
        ],
        col_widths=[3, 6.5, 6.5],
    )
    add_para(doc, "（2）软件环境", bold=True, size=12, space_after=4)
    add_bullet(doc, "操作系统：Windows 10/11 x64；DFTB+ 计算运行于 WSL2 Ubuntu 发行版。")
    add_bullet(doc, "学生端安装包内嵌 Windows Python 与依赖，用户无需另行预装 Python。")
    add_bullet(doc, "首次使用需完成“部署到本机”（含 micromamba 环境、SK 参数与冒烟校验）。")
    add_bullet(doc, "可选：配置 DeepSeek API，以增强自然语言生成 HSD 的能力。")

    add_page_break(doc)

    # ========== 2 ==========
    add_heading_cn(doc, "2 安装与启动", 1)
    add_heading_cn(doc, "2.1 学生端安装", 2)
    add_body(doc, "学生端安装包文件名形如 DFTB_Neu_Setup_*.exe。双击运行安装程序，按界面提示选择安装目录并完成安装。")
    add_bullet(doc, "安装完成后，可在开始菜单或桌面快捷方式中找到“DFTB 工作台 / DFTB-NEU”。")
    add_bullet(doc, "安装包为现代 Setup（Electron 界面），安装后即可打开使用。")
    add_bullet(doc, "计算环境并不在安装瞬间全部就绪，需在软件内执行“部署到本机”。")

    add_heading_cn(doc, "2.2 教师端安装", 2)
    add_body(
        doc,
        "教师端安装包文件名形如 DFTB_Neu_Teacher_Setup_*.exe。"
        "安装后打开教师端即可使用；课堂服务地址与教师凭证由安装包预置，日常一般无需在界面手工填写。"
        "教师端仅负责课堂门禁与会话监管，不执行 DFTB+ 计算。",
    )

    add_heading_cn(doc, "2.3 启动与登录", 2)
    add_body(doc, "启动学生端后进入登录页，按下列步骤登录：")
    add_bullet(doc, "输入学号（或教师指定的用户标识）。")
    add_bullet(doc, "输入本堂密码（由教师端发布并口头/群公告知）。")
    add_bullet(doc, "点击“进入工作台”。")
    add_body(
        doc,
        "口令可为教师当堂发布的短密码，也可为运维配置的永久凭证（长期会话）。"
        "若无法登录，请核对本堂密码是否仍有效，并确认课堂中心服务可访问。",
    )
    add_image(
        doc,
        PPT_MEDIA / "image1.png",
        width_cm=14.5,
        caption="图 2-1  软件整体界面与产品外观示意（登录/工作台主视觉）",
    )

    add_page_break(doc)

    # ========== 3 ==========
    add_heading_cn(doc, "3 学生端功能说明", 1)
    add_heading_cn(doc, "3.1 界面总览", 2)
    add_body(
        doc,
        "学生端主界面顶部提供功能导航，主要包括：首页、计算、任务、课程、环境、设置。"
        "用户可按“登录 → 环境部署 → 课程训练/自由计算 → 任务查看”的顺序完成一次完整上课或科研计算流程。",
    )
    add_table(
        doc,
        ["导航页", "功能摘要"],
        [
            ["首页", "一句话计算入口、示例指令、快速上手指南与核心能力介绍"],
            ["环境", "查看 WSL / DFTB+ / 冒烟校验状态，执行“部署到本机”"],
            ["课程", "按固体物理课次载入结构与计算意图"],
            ["计算", "自然语言对话生成 HSD，上传/编辑结构，确认后投递"],
            ["任务", "跟踪任务进度，查看能量/能带/DOS/轨迹等结果并导出"],
            ["设置", "配色主题、DeepSeek API 等本地配置"],
        ],
        col_widths=[3, 13],
    )

    add_heading_cn(doc, "3.2 首页", 2)
    add_body(
        doc,
        "首页以“用一句话完成固体计算”为核心交互：用户在自然语言计算指令框中输入中文描述，"
        "也可点击预置示例按钮快速填入指令，例如“硅 · 能带/DOS”“石墨烯 · 能带”“水 · 几何优化”等。"
        "点击“开始计算”后进入计算流程；亦可打开完整计算台进行更细粒度控制。",
    )
    add_image(
        doc,
        PPT_MEDIA / "image3.png",
        width_cm=14.5,
        caption="图 3-1  学生端首页与一句话计算入口示意",
    )

    add_heading_cn(doc, "3.3 环境部署", 2)
    add_body(
        doc,
        "“环境”页用于完成本机 DFTB+ 计算环境准备。界面展示三项关键状态：WSL、DFTB+、冒烟校验。"
        "三项均通过后，方可稳定开展课程与自由计算。",
    )
    add_para(doc, "操作步骤：", bold=True, size=12, space_after=4)
    add_bullet(doc, "打开“环境”页，点击“刷新”查看当前状态。")
    add_bullet(doc, "若未就绪，点击“部署到本机”。部署将写入用户目录 ~/.dftb-neu（隔离环境，不修改系统全局配置）。")
    add_bullet(doc, "首次可能需要管理员确认启用 WSL2；若系统提示重启，请重启后再次点击“部署到本机”。")
    add_bullet(doc, "部署包含 micromamba 环境、SK 参数包，以及 H₂O 几何优化冒烟校验。")
    add_body(
        doc,
        "页面同时提供分步部署教程：启用 WSL、使用国内镜像下载 Ubuntu rootfs、导入发行版、创建日常用户、"
        "回到软件完成 DFTB+ 部署等。推荐优先使用软件内一键部署流程；仅在网络或权限受限时按教程手工操作。",
    )
    add_image(
        doc,
        PPT_MEDIA / "image4.png",
        width_cm=14.5,
        caption="图 3-2  环境页与“部署到本机”示意",
    )

    add_page_break(doc)

    add_heading_cn(doc, "3.4 固体物理课程", 2)
    add_body(
        doc,
        "“课程”页预置固体物理课次，对齐本科固体物理教学主线（晶体结构 → 结合与能量 → 晶格振动 → 能带 → 二维与缺陷等）。"
        "学生按课次载入结构文件与计算意图，聚焦物理图像理解，而不是软件配置细节。",
    )
    add_table(
        doc,
        ["序号", "课次主题", "典型计算"],
        [
            ["1", "晶体结构与倒格矢", "Si SCC 单点"],
            ["2", "固体结合与总能量", "能量与 Fermi 能级"],
            ["3", "晶格振动：平衡结构", "几何优化"],
            ["4", "能带理论：Si 能带/DOS", "Band + DOS"],
            ["5", "二维晶体：石墨烯", "Dirac 锥图像"],
            ["6", "缺陷与局域态", "石墨烯空位"],
            ["7", "分子动力学入门", "BOMD"],
            ["8", "科研综合课题", "自选结构 + 方法摘要"],
        ],
        col_widths=[2, 7, 7],
    )
    add_image(
        doc,
        PPT_MEDIA / "image7.png",
        width_cm=14.5,
        caption="图 3-3  固体物理课程页示意",
    )

    add_heading_cn(doc, "3.5 自然语言计算", 2)
    add_body(
        doc,
        "“计算”页是科研与拓展训练的核心界面。用户可在对话区用中文描述任务，点击“生成输入”由 AI 大模型辅助生成 DFTB+ HSD；"
        "也可“恢复推荐默认”。结构可通过上传 POSCAR/GEN、粘贴结构文本，或从课程/Materials Project 获取。",
    )
    add_para(doc, "关键规则：", bold=True, size=12, space_after=4)
    add_bullet(doc, "周期体系（能带、缺陷、声子等）必须提供 POSCAR/GEN。")
    add_bullet(doc, "分子示例（几何优化、Casida、BOMD 等）在无结构时可使用内置水分子示例。")
    add_bullet(doc, "生成的 HSD 必须经用户在界面确认后才会本机投递，AI 不自动代跑。")
    add_bullet(doc, "确认计算后，任务进入本机队列，可在“任务”页跟踪。")
    add_image(
        doc,
        PPT_MEDIA / "image6.png",
        width_cm=14.5,
        caption="图 3-4  计算页（自然语言 → HSD → 确认投递）示意",
    )

    add_page_break(doc)

    add_heading_cn(doc, "3.6 任务与结果", 2)
    add_body(
        doc,
        "“任务”页用于查看已投递计算的进度与结果。用户可在任务列表中切换“上一个/下一个”，取消未完成任务，"
        "下载输出文件，并查看结构对比、结果图（如能带/DOS、轨迹等）。支持导出图件与方法摘要。",
    )
    add_bullet(doc, "可识读关键观测量：总能量、Fermi 能级、能带/DOS、几何优化收敛信息、BOMD 轨迹等。")
    add_bullet(doc, "“导出方法摘要”生成方法段与课堂摘要（Markdown/HTML），用于报告或作业，并非完整论文。")
    add_bullet(doc, "导出内容建议核验 SK 参数族与能力表述是否与实际计算一致。")
    add_image(
        doc,
        PPT_MEDIA / "image5.png",
        width_cm=14.5,
        caption="图 3-5  任务页与结果可视化示意",
    )

    add_heading_cn(doc, "3.7 设置", 2)
    add_body(doc, "“设置”页提供个性化与外部服务配置：")
    add_bullet(doc, "软件系统配色：切换界面主题。")
    add_bullet(doc, "DeepSeek API：填写 API Key、Base URL 与模型名称，点击“保存并检测”。")
    add_bullet(doc, "课堂服务地址一般由安装包预置；日常使用主要依赖学号与本堂密码。")
    add_body(doc, "说明：API Key 属于个人密钥，请妥善保管，勿写入公开文档或截图外传。")

    add_page_break(doc)

    # ========== 4 ==========
    add_heading_cn(doc, "4 教师端功能说明", 1)
    add_body(
        doc,
        "教师端界面主题为“课堂门禁与会话”，用于课前发布口令、课中监管学生、课后结束本堂。"
        "教师端不承担 DFTB+ 计算，计算由学生在各自本机完成。",
    )
    add_heading_cn(doc, "4.1 本堂密码管理", 2)
    add_bullet(doc, "点击“发布本堂密码”生成/添加当堂短密码，并告知学生。")
    add_bullet(doc, "可叠加多个本堂短密码；列表中可复制或单独吊销。")
    add_bullet(doc, "“清空全部”将清空全部短密码并注销对应短密码会话。")
    add_bullet(doc, "永久密码用于长期会话；修改永久密码需输入新密码并确认。")

    add_heading_cn(doc, "4.2 在线学生监管", 2)
    add_body(
        doc,
        "“已登录学生”列表自动刷新，展示学号、活跃状态、登录/过期时间、今日 LLM 调用情况、设备信息等，"
        "并支持对异常会话执行注销。结课时点击“结束本堂”，清空口令并使短密码会话失效。",
    )

    add_heading_cn(doc, "4.3 建议课时流程", 2)
    add_table(
        doc,
        ["阶段", "建议用时", "教师动作", "学生动作"],
        [
            ["环境", "约 15 分钟", "讲解三项状态含义", "完成“部署到本机”"],
            ["课程训练", "主课时", "按课次引导物理图像", "载入课次并完成计算"],
            ["自由计算", "拓展", "巡视确认 HSD", "自然语言描述并确认投递"],
            ["成稿", "课后/尾声", "检查方法表述", "导出方法摘要"],
        ],
        col_widths=[3, 3, 5, 5],
    )

    add_page_break(doc)

    # ========== 5 ==========
    add_heading_cn(doc, "5 典型操作流程", 1)
    add_heading_cn(doc, "5.1 课堂教学完整流程", 2)
    add_bullet(doc, "教师确认课堂中心服务可用，打开教师端并发布本堂密码。")
    add_bullet(doc, "学生安装并启动学生端，使用学号 + 本堂密码登录。")
    add_bullet(doc, "学生在“环境”页执行“部署到本机”，直至三项状态通过。")
    add_bullet(doc, "在“课程”页按课次顺序载入结构与意图，完成课堂计算。")
    add_bullet(doc, "需要拓展时，在“计算”页用自然语言描述任务，确认 HSD 后投递。")
    add_bullet(doc, "在“任务”页查看结果，导出图件或方法摘要。")
    add_bullet(doc, "下课前教师点击“结束本堂”，清空口令。")

    add_heading_cn(doc, "5.2 科研自由计算流程", 2)
    add_bullet(doc, "准备或上传 POSCAR/GEN（周期体系必需）。")
    add_bullet(doc, "在计算页用自然语言描述目标（如几何优化、能带/DOS、BOMD、激发态等）。")
    add_bullet(doc, "生成并检查 HSD：核对 Driver、Hamiltonian、SK 参数与 k 点等关键项。")
    add_bullet(doc, "确认计算，等待本机任务完成。")
    add_bullet(doc, "下载输出、导出观测量图件，必要时导出方法摘要用于记录可复现流程。")

    add_heading_cn(doc, "5.3 推荐评分观察点（教学）", 2)
    add_table(
        doc,
        ["评分项", "观察点"],
        [
            ["输入", "HSD 是否与意图一致（Driver / Hamiltonian）"],
            ["计算", "冒烟或作业完成，日志无致命错误"],
            ["分析", "能指认总能量 / 能带或轨迹相关输出"],
            ["文稿", "方法摘要含 DFTB+、SK、能力族等关键信息"],
        ],
        col_widths=[3, 13],
    )

    add_page_break(doc)

    # ========== 6 ==========
    add_heading_cn(doc, "6 常见问题与故障排查", 1)
    add_table(
        doc,
        ["现象", "可能原因", "处理建议"],
        [
            ["软件无法启动", "后端进程异常或端口占用", "查看 %APPDATA%\\DFTB工作台\\backend.log，重启应用"],
            ["无法登录", "本堂密码错误/已吊销，或课堂服务不可达", "向教师核对本堂密码；检查网络与课堂中心"],
            ["部署失败", "WSL 未启用/需重启/镜像下载失败", "按环境页教程启用 WSL，优先国内镜像，重启后重试"],
            ["计算无法投递", "未确认 HSD，或结构缺失", "周期体系补齐 POSCAR/GEN；确认后再投递"],
            ["AI 生成异常", "API Key 无效或网络受限", "在设置页检测 DeepSeek API，核对 Key 与模型"],
            ["结果图缺失", "任务未完成或输出未生成", "在任务页查看日志与状态，等待完成后再导出"],
        ],
        col_widths=[3.5, 5.5, 7],
    )
    add_body(
        doc,
        "数据与配置位置说明：WSL 内数据根目录为 $HOME/.dftb-neu；Windows 侧配置通常位于 %LOCALAPPDATA%\\dftb-neu。"
        "排查问题时可结合环境页日志与上述目录中的运行信息。",
    )

    add_page_break(doc)

    # ========== 7 ==========
    add_heading_cn(doc, "7 注意事项与免责声明", 1)
    add_heading_cn(doc, "7.1 使用注意", 2)
    add_bullet(doc, "本软件用于教学与科研辅助，计算结果需结合专业判断审慎使用。")
    add_bullet(doc, "AI 生成的输入文件必须人工确认；错误参数可能导致计算失败或无意义结果。")
    add_bullet(doc, "请勿在公开场合泄露 API Key、教师永久凭证等敏感信息。")
    add_bullet(doc, "部署环境写入用户目录，一般不改系统全局；但仍建议在个人教学/科研电脑上使用。")
    add_bullet(doc, "课堂短密码具有时效性，结课后应及时结束本堂，避免口令残留。")

    add_heading_cn(doc, "7.2 知识产权说明", 2)
    add_body(
        doc,
        f"本手册所述软件名称为“{SOFT_FULL}”，版本 {SOFT_VER}。"
        "软件著作权归东北大学所有。未经权利人许可，不得擅自复制、传播、反向工程或以其他方式侵害本软件知识产权。"
        "第三方组件（如 DFTB+、相关科学计算库、前端依赖等）遵循其各自开源或商业许可协议。",
    )

    add_heading_cn(doc, "7.3 文档修订", 2)
    add_table(
        doc,
        ["版本", "日期", "说明"],
        [
            ["V1.0", "2026-08", "首版用户手册，用于软件著作权登记申请"],
        ],
        col_widths=[3, 4, 9],
    )

    add_para(doc, "", space_after=18)
    add_para(doc, "—— 手册结束 ——", size=12, align="center", space_after=6)
    add_para(doc, HEADER, size=10.5, align="center", space_after=0)

    out = OUT_DIR / f"{SOFT_FULL}_{SOFT_VER}_用户手册.docx"
    doc.save(str(out))
    print(out)
    return out


if __name__ == "__main__":
    build()
