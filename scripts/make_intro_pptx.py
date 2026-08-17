"""Generate DFTB-NEU product introduction PowerPoint."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Brand colors from product UI
ACCENT = RGBColor(0x0F, 0x6E, 0x6A)
ACCENT_DIM = RGBColor(0x0B, 0x56, 0x53)
ACCENT_HI = RGBColor(0x2A, 0x8F, 0x7F)
INK = RGBColor(0x1C, 0x24, 0x30)
MUTED = RGBColor(0x5A, 0x66, 0x75)
BG = RGBColor(0xF3, 0xF5, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MARK_FROM = RGBColor(0x0A, 0x3D, 0x3B)
LINE = RGBColor(0xD0, 0xD5, 0xDB)
SOFT = RGBColor(0xA8, 0xC5, 0xC2)
SOFT2 = RGBColor(0xC8, 0xDD, 0xDB)
SOFT3 = RGBColor(0x9A, 0xB8, 0xB5)
SOFT4 = RGBColor(0xB8, 0xD4, 0xD1)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W = prs.slide_width
H = prs.slide_height
blank = prs.slide_layouts[6]
TOTAL = 10


def add_rect(slide, left, top, width, height, fill_rgb):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    shape.line.fill.background()
    return shape


def set_run(run, text, size=18, bold=False, color=INK, font_name="Microsoft YaHei"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return run


def add_textbox(slide, left, top, width, height, paragraphs, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    anchor_map = {
        MSO_ANCHOR.TOP: "t",
        MSO_ANCHOR.MIDDLE: "ctr",
        MSO_ANCHOR.BOTTOM: "b",
    }
    try:
        tf._txBody.bodyPr.set("anchor", anchor_map.get(valign, "t"))
    except Exception:
        pass
    for i, p in enumerate(paragraphs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if "space_after" in p:
            para.space_after = Pt(p["space_after"])
        if "space_before" in p:
            para.space_before = Pt(p["space_before"])
        run = para.add_run()
        set_run(
            run,
            p["text"],
            size=p.get("size", 18),
            bold=p.get("bold", False),
            color=p.get("color", INK),
            font_name=p.get("font", "Microsoft YaHei"),
        )
    return box


def add_accent_bar(slide):
    return add_rect(slide, Inches(0.7), Inches(0.55), Inches(0.12), Inches(0.45), ACCENT)


def slide_chrome(slide, page_num, total=TOTAL, dark=False):
    if dark:
        add_rect(slide, 0, H - Inches(0.42), W, Inches(0.42), MARK_FROM)
        add_textbox(
            slide,
            Inches(0.7),
            H - Inches(0.4),
            Inches(8),
            Inches(0.36),
            [
                {
                    "text": "DFTB-NEU  ·  东北大学材料科学与工程学院",
                    "size": 11,
                    "color": SOFT,
                }
            ],
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_textbox(
            slide,
            W - Inches(1.6),
            H - Inches(0.4),
            Inches(1.2),
            Inches(0.36),
            [
                {
                    "text": f"{page_num} / {total}",
                    "size": 11,
                    "color": SOFT,
                    "align": PP_ALIGN.RIGHT,
                }
            ],
            valign=MSO_ANCHOR.MIDDLE,
        )
    else:
        add_rect(slide, 0, H - Inches(0.42), W, Inches(0.42), BG)
        add_rect(slide, 0, H - Inches(0.42), W, Pt(1), LINE)
        add_textbox(
            slide,
            Inches(0.7),
            H - Inches(0.4),
            Inches(9),
            Inches(0.36),
            [
                {
                    "text": "DFTB-NEU  ·  东北大学材料科学与工程学院 · 张林教授课题组",
                    "size": 11,
                    "color": MUTED,
                }
            ],
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_textbox(
            slide,
            W - Inches(1.6),
            H - Inches(0.4),
            Inches(1.2),
            Inches(0.36),
            [
                {
                    "text": f"{page_num} / {total}",
                    "size": 11,
                    "color": MUTED,
                    "align": PP_ALIGN.RIGHT,
                }
            ],
            valign=MSO_ANCHOR.MIDDLE,
        )


# ========== SLIDE 1: Cover ==========
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, W, H, MARK_FROM)
add_rect(s, 0, 0, Inches(0.18), H, ACCENT_HI)
add_rect(s, W - Inches(4.2), 0, Inches(4.2), H, ACCENT_DIM)
add_textbox(
    s,
    Inches(0.9),
    Inches(1.6),
    Inches(7.5),
    Inches(0.5),
    [{"text": "DFTB+  ·  Local Workspace", "size": 16, "color": ACCENT_HI, "bold": True}],
)
add_textbox(
    s,
    Inches(0.9),
    Inches(2.15),
    Inches(7.8),
    Inches(1.2),
    [{"text": "DFTB-NEU 工作台", "size": 48, "bold": True, "color": WHITE, "space_after": 8}],
)
add_textbox(
    s,
    Inches(0.9),
    Inches(3.4),
    Inches(7.5),
    Inches(1.0),
    [
        {"text": "面向科研与教学的 DFTB+ 本机计算工作台", "size": 22, "color": SOFT2},
        {
            "text": "Windows 桌面壳  ·  本机 API  ·  WSL2 隔离环境",
            "size": 16,
            "color": SOFT3,
            "space_before": 10,
        },
    ],
)
add_textbox(
    s,
    Inches(0.9),
    Inches(5.5),
    Inches(7.5),
    Inches(0.8),
    [
        {"text": "东北大学 · 材料科学与工程学院", "size": 15, "color": WHITE},
        {"text": "张林教授课题组", "size": 14, "color": ACCENT_HI, "space_before": 4},
    ],
)
add_textbox(
    s,
    W - Inches(3.7),
    Inches(2.8),
    Inches(3.2),
    Inches(2.2),
    [
        {"text": "D+", "size": 72, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER},
        {
            "text": "产品介绍",
            "size": 18,
            "color": ACCENT_HI,
            "align": PP_ALIGN.CENTER,
            "space_before": 12,
        },
    ],
)

# ========== SLIDE 2: Agenda ==========
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, W, H, WHITE)
add_accent_bar(s)
add_textbox(
    s,
    Inches(1.0),
    Inches(0.5),
    Inches(10),
    Inches(0.55),
    [{"text": "目录", "size": 32, "bold": True, "color": INK}],
)
items = [
    ("01", "产品定位", "是什么、解决什么问题"),
    ("02", "为什么需要", "教学与科研场景痛点"),
    ("03", "系统架构", "三端协同与本机计算"),
    ("04", "核心能力", "部署 · 课程 · 自然语言计算"),
    ("05", "固体物理课次", "对齐黄昆体系的课堂演示"),
    ("06", "课时流程", "从环境到计算"),
    ("07", "技术亮点", "隔离部署与本机计算"),
    ("08", "总结与展望", "一句话带走"),
]
for i, (num, title, sub) in enumerate(items):
    col = i // 4
    row = i % 4
    left = Inches(0.9) + Inches(6.0) * col
    top = Inches(1.4) + Inches(1.15) * row
    add_textbox(
        s,
        left,
        top,
        Inches(0.7),
        Inches(0.5),
        [{"text": num, "size": 22, "bold": True, "color": ACCENT}],
    )
    add_textbox(
        s,
        left + Inches(0.85),
        top,
        Inches(4.5),
        Inches(0.85),
        [
            {"text": title, "size": 20, "bold": True, "color": INK, "space_after": 4},
            {"text": sub, "size": 13, "color": MUTED},
        ],
    )
slide_chrome(s, 2)

# ========== SLIDE 3: Positioning ==========
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, W, H, WHITE)
add_accent_bar(s)
add_textbox(
    s,
    Inches(1.0),
    Inches(0.5),
    Inches(11),
    Inches(0.55),
    [{"text": "产品定位", "size": 32, "bold": True, "color": INK}],
)
add_textbox(
    s,
    Inches(0.9),
    Inches(1.25),
    Inches(11.5),
    Inches(0.6),
    [
        {
            "text": "DFTB+ 本机计算工作台：把高水平原子尺度计算，装进课堂与课题组日常。",
            "size": 18,
            "color": MUTED,
        }
    ],
)
cards = [
    ("本机优先", "计算在学生本机 WSL2 中完成\n不依赖远程超算即可上课"),
    ("教学就绪", "预置固体物理课次与结构\n教师端发布本堂密码监管"),
    ("科研可用", "自然语言描述任务 → HSD\n本机投递，服务课题组日常"),
]
for i, (t, body) in enumerate(cards):
    left = Inches(0.9) + Inches(4.0) * i
    add_rect(s, left, Inches(2.2), Inches(3.7), Inches(3.6), BG)
    add_rect(s, left, Inches(2.2), Inches(3.7), Inches(0.1), ACCENT)
    add_textbox(
        s,
        left + Inches(0.3),
        Inches(2.55),
        Inches(3.1),
        Inches(0.5),
        [{"text": t, "size": 20, "bold": True, "color": ACCENT}],
    )
    add_textbox(
        s,
        left + Inches(0.3),
        Inches(3.3),
        Inches(3.1),
        Inches(2.2),
        [{"text": body, "size": 15, "color": INK}],
    )
slide_chrome(s, 3)

# ========== SLIDE 4: Why ==========
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, W, H, WHITE)
add_accent_bar(s)
add_textbox(
    s,
    Inches(1.0),
    Inches(0.5),
    Inches(11),
    Inches(0.55),
    [{"text": "为什么需要 DFTB-NEU", "size": 32, "bold": True, "color": INK}],
)
pain = [
    ("环境门槛高", "DFTB+、SK 参数、WSL 配置复杂，课堂难统一就绪"),
    ("输入门槛高", "HSD / POSCAR 对本科生不友好，演示时间被配置吃掉"),
    ("教学难闭环", "缺少课次编排与课堂在线监管，组织成本高"),
    ("科研衔接弱", "课堂与课题组工具割裂，同一环境难以延续自由课题"),
]
for i, (t, d) in enumerate(pain):
    top = Inches(1.35) + Inches(1.2) * i
    add_rect(s, Inches(0.9), top, Inches(0.12), Inches(0.9), ACCENT)
    add_textbox(
        s,
        Inches(1.3),
        top,
        Inches(3.2),
        Inches(0.9),
        [{"text": t, "size": 18, "bold": True, "color": INK}],
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        s,
        Inches(4.7),
        top,
        Inches(7.8),
        Inches(0.9),
        [{"text": d, "size": 16, "color": MUTED}],
        valign=MSO_ANCHOR.MIDDLE,
    )
slide_chrome(s, 4)

# ========== SLIDE 5: Architecture ==========
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, W, H, WHITE)
add_accent_bar(s)
add_textbox(
    s,
    Inches(1.0),
    Inches(0.5),
    Inches(11),
    Inches(0.55),
    [{"text": "系统架构 · 三端协同", "size": 32, "bold": True, "color": INK}],
)
roles = [
    (
        "学生端",
        "DFTB-NEU 工作台",
        [
            "学号 + 本堂密码登录",
            "部署到本机（WSL / DFTB+ / 冒烟）",
            "课程 · 计算 · 任务与观测量",
        ],
        ACCENT,
    ),
    (
        "教师端",
        "DFTB-NEU 教师端",
        [
            "发布 / 吊销本堂密码",
            "监管在线学生与会话",
            "结束本堂（无计算职责）",
        ],
        ACCENT_DIM,
    ),
    (
        "课堂中心",
        "Classroom Hub",
        [
            "鉴权与会话管理",
            "LLM / Materials Project 代理",
            "内网服务，安装包预置地址",
        ],
        MARK_FROM,
    ),
]
for i, (role, name, bullets, color) in enumerate(roles):
    left = Inches(0.7) + Inches(4.15) * i
    add_rect(s, left, Inches(1.35), Inches(3.9), Inches(5.0), BG)
    add_rect(s, left, Inches(1.35), Inches(3.9), Inches(1.15), color)
    add_textbox(
        s,
        left + Inches(0.25),
        Inches(1.45),
        Inches(3.4),
        Inches(0.95),
        [
            {"text": role, "size": 14, "color": SOFT4, "space_after": 2},
            {"text": name, "size": 18, "bold": True, "color": WHITE},
        ],
    )
    paras = [
        {"text": "·  " + b, "size": 15, "color": INK, "space_before": 14 if j else 0}
        for j, b in enumerate(bullets)
    ]
    add_textbox(s, left + Inches(0.3), Inches(2.8), Inches(3.3), Inches(3.2), paras)
slide_chrome(s, 5)

# ========== SLIDE 6: Core capabilities ==========
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, W, H, WHITE)
add_accent_bar(s)
add_textbox(
    s,
    Inches(1.0),
    Inches(0.5),
    Inches(11),
    Inches(0.55),
    [{"text": "核心能力", "size": 32, "bold": True, "color": INK}],
)
caps = [
    (
        "01  环境部署",
        "「部署到本机」：micromamba 环境、SK 参数、H₂O 冒烟校验；写入 ~/.dftb-neu，不改系统全局。",
    ),
    (
        "02  课程训练",
        "预置固体物理课次与结构文件；按课载入意图，聚焦物理图像而非软件配置。",
    ),
    (
        "03  自然语言计算",
        "用中文描述任务 → 确认 HSD → 本机投递；支持能带、几何优化、BOMD、Casida 等。",
    ),
    (
        "04  结构来源",
        "课程内置结构，或从 Materials Project 拉取 POSCAR/GEN；周期体系与分子场景均可。",
    ),
    (
        "05  任务与观测量",
        "任务页跟踪进度，识读总能量、能带/DOS、轨迹等关键输出。",
    ),
    (
        "06  教学登录与监管",
        "学号 + 本堂密码登录；教师端发布口令、查看在线会话，结课一键清空。",
    ),
]
for i, (t, d) in enumerate(caps):
    col = i % 2
    row = i // 2
    left = Inches(0.7) + Inches(6.25) * col
    top = Inches(1.3) + Inches(1.7) * row
    add_rect(s, left, top, Inches(5.95), Inches(1.5), BG)
    add_textbox(
        s,
        left + Inches(0.25),
        top + Inches(0.2),
        Inches(5.45),
        Inches(1.15),
        [
            {"text": t, "size": 16, "bold": True, "color": ACCENT, "space_after": 6},
            {"text": d, "size": 13, "color": INK},
        ],
    )
slide_chrome(s, 6)

# ========== SLIDE 7: Curriculum ==========
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, W, H, WHITE)
add_accent_bar(s)
add_textbox(
    s,
    Inches(1.0),
    Inches(0.5),
    Inches(11),
    Inches(0.55),
    [{"text": "固体物理课次（预置）", "size": 32, "bold": True, "color": INK}],
)
add_textbox(
    s,
    Inches(0.9),
    Inches(1.15),
    Inches(11.5),
    Inches(0.45),
    [
        {
            "text": "对齐黄昆《固体物理学》本科主线：晶体结构 → 结合与能量 → 晶格振动 → 能带 → 二维与缺陷",
            "size": 14,
            "color": MUTED,
        }
    ],
)
lessons = [
    ("01", "晶体结构与倒格矢", "Si SCC 单点"),
    ("02", "固体结合与总能量", "能量与 Fermi"),
    ("03", "晶格振动：平衡结构", "几何优化"),
    ("04", "能带理论：Si 能带/DOS", "Band + DOS"),
    ("05", "二维晶体：石墨烯", "Dirac 锥图像"),
    ("06", "缺陷与局域态", "石墨烯空位"),
]
for i, (num, title, tag) in enumerate(lessons):
    col = i % 3
    row = i // 3
    left = Inches(0.7) + Inches(4.15) * col
    top = Inches(1.85) + Inches(2.15) * row
    add_rect(s, left, top, Inches(3.95), Inches(1.9), BG)
    add_rect(s, left, top, Inches(0.1), Inches(1.9), ACCENT)
    add_textbox(
        s,
        left + Inches(0.35),
        top + Inches(0.35),
        Inches(3.3),
        Inches(1.3),
        [
            {"text": num, "size": 14, "bold": True, "color": ACCENT, "space_after": 6},
            {"text": title, "size": 17, "bold": True, "color": INK, "space_after": 8},
            {"text": tag, "size": 13, "color": MUTED},
        ],
    )
slide_chrome(s, 7)

# ========== SLIDE 8: Workflow ==========
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, W, H, WHITE)
add_accent_bar(s)
add_textbox(
    s,
    Inches(1.0),
    Inches(0.5),
    Inches(11),
    Inches(0.55),
    [{"text": "建议课时流程", "size": 32, "bold": True, "color": INK}],
)
steps = [
    ("1", "环境", "15 min", "解读 WSL / DFTB+ / 冒烟\n三项状态，完成部署"),
    ("2", "课程训练", "主课时", "按课次顺序载入结构\n与计算意图，边算边讲"),
    ("3", "自由计算", "拓展", "自然语言描述任务\n确认 HSD 后投递本机"),
]
for i, (n, title, time, body) in enumerate(steps):
    left = Inches(1.4) + Inches(3.7) * i
    add_rect(s, left + Inches(1.25), Inches(1.5), Inches(0.7), Inches(0.7), ACCENT)
    add_textbox(
        s,
        left + Inches(1.25),
        Inches(1.5),
        Inches(0.7),
        Inches(0.7),
        [{"text": n, "size": 22, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}],
        valign=MSO_ANCHOR.MIDDLE,
    )
    if i < 2:
        add_rect(s, left + Inches(2.05), Inches(1.8), Inches(2.7), Pt(3), ACCENT_HI)
    add_textbox(
        s,
        left + Inches(0.2),
        Inches(2.5),
        Inches(3.1),
        Inches(3.5),
        [
            {
                "text": title,
                "size": 20,
                "bold": True,
                "color": INK,
                "align": PP_ALIGN.CENTER,
                "space_after": 6,
            },
            {
                "text": time,
                "size": 13,
                "color": ACCENT,
                "align": PP_ALIGN.CENTER,
                "space_after": 14,
            },
            {"text": body, "size": 14, "color": MUTED, "align": PP_ALIGN.CENTER},
        ],
    )
slide_chrome(s, 8)

# ========== SLIDE 9: Tech highlights ==========
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, W, H, WHITE)
add_accent_bar(s)
add_textbox(
    s,
    Inches(1.0),
    Inches(0.5),
    Inches(11),
    Inches(0.55),
    [{"text": "技术亮点", "size": 32, "bold": True, "color": INK}],
)
techs = [
    (
        "Windows + WSL2",
        "桌面壳（Electron）内嵌本机 API；DFTB+ 在隔离发行版中运行，首次部署含冒烟校验。",
    ),
    (
        "零全局污染",
        "环境写入用户目录 ~/.dftb-neu；安装包内嵌 Python，学生无需预装解释器。",
    ),
    (
        "教学鉴权模型",
        "本堂短密码可叠加/吊销；永久凭证保留长期会话；结课一键清空口令。",
    ),
    (
        "AI 辅助但不替代",
        "自然语言 → HSD 需用户确认后再投递，降低误算；计算在本机完成，过程可追溯。",
    ),
]
for i, (t, d) in enumerate(techs):
    top = Inches(1.35) + Inches(1.25) * i
    add_rect(s, Inches(0.9), top, Inches(11.5), Inches(1.1), BG)
    add_rect(s, Inches(0.9), top, Inches(0.12), Inches(1.1), ACCENT)
    add_textbox(
        s,
        Inches(1.3),
        top + Inches(0.2),
        Inches(10.8),
        Inches(0.75),
        [
            {"text": t, "size": 17, "bold": True, "color": ACCENT, "space_after": 4},
            {"text": d, "size": 14, "color": INK},
        ],
    )
slide_chrome(s, 9)

# ========== SLIDE 10: Summary ==========
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, W, H, MARK_FROM)
add_rect(s, 0, 0, Inches(0.18), H, ACCENT_HI)
add_textbox(
    s,
    Inches(0.9),
    Inches(1.5),
    Inches(11.5),
    Inches(0.5),
    [{"text": "一句话总结", "size": 16, "color": ACCENT_HI, "bold": True}],
)
add_textbox(
    s,
    Inches(0.9),
    Inches(2.1),
    Inches(11.5),
    Inches(1.8),
    [
        {
            "text": "DFTB-NEU 让固体物理课堂\n真正「算得起来、讲得清楚」。",
            "size": 28,
            "bold": True,
            "color": WHITE,
        }
    ],
)
add_textbox(
    s,
    Inches(0.9),
    Inches(4.3),
    Inches(11.5),
    Inches(1.2),
    [
        {
            "text": "本机隔离部署  ·  课次即用  ·  自然语言到 HSD  ·  三端协同教学",
            "size": 16,
            "color": SOFT,
        },
        {"text": "", "size": 10},
        {
            "text": "东北大学 · 材料科学与工程学院 · 张林教授课题组",
            "size": 14,
            "color": WHITE,
            "space_before": 16,
        },
    ],
)
add_textbox(
    s,
    Inches(0.9),
    Inches(6.5),
    Inches(11.5),
    Inches(0.4),
    [{"text": "感谢聆听  ·  欢迎交流", "size": 14, "color": ACCENT_HI}],
)

out = "docs/DFTB-NEU产品介绍.pptx"
try:
    prs.save(out)
except PermissionError:
    out = "docs/DFTB-NEU产品介绍_无成稿.pptx"
    prs.save(out)
print("saved", out)
