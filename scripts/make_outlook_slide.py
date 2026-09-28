"""单页：总结与展望 · 未来优化方向（对齐 DFTB-NEU 产品介绍 PPT 风格）"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ACCENT = RGBColor(0x0F, 0x6E, 0x6A)
INK = RGBColor(0x1C, 0x24, 0x30)
MUTED = RGBColor(0x5A, 0x66, 0x75)
BG = RGBColor(0xF3, 0xF5, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD0, 0xD5, 0xDB)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]


def add_rect(slide, left, top, width, height, fill_rgb):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    shape.line.fill.background()
    return shape


def set_run(run, text, size=18, bold=False, color=INK, font_name="Microsoft YaHei"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name


def add_textbox(slide, left, top, width, height, paragraphs, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(
                valign, "t"
            ),
        )
    except Exception:
        pass
    for i, p in enumerate(paragraphs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if "space_after" in p:
            para.space_after = Pt(p["space_after"])
        if "space_before" in p:
            para.space_before = Pt(p["space_before"])
        run = para.add_run()
        set_run(
            run,
            p["text"],
            size=p.get("size", 18),
            bold=p.get("bold", False),
            color=p.get("color", INK),
            font_name=p.get("font", "Microsoft YaHei"),
        )
    return box


s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, W, H, WHITE)
add_rect(s, Inches(0.7), Inches(0.55), Inches(0.12), Inches(0.45), ACCENT)

add_textbox(
    s,
    Inches(1.0),
    Inches(0.48),
    Inches(10),
    Inches(0.55),
    [{"text": "总结与展望", "size": 32, "bold": True, "color": INK}],
)
add_textbox(
    s,
    Inches(1.0),
    Inches(1.1),
    Inches(11.5),
    Inches(0.45),
    [
        {
            "text": "未来优化方向  ·  让课堂计算更贴近真实科研",
            "size": 16,
            "color": MUTED,
        }
    ],
)

cards = [
    (
        "01",
        "贴近真实科研",
        "从课堂演示走向可复现科研工作流",
        "结构 → 计算 → 分析 → 报告一气呵成；参数可追溯、结果可对比；支持课题级数据与长期项目沉淀。",
    ),
    (
        "02",
        "完善多功能",
        "补齐教—学—管闭环与计算场景",
        "布置、提交、批阅、复现、监管一体；扩展能带/DOS、批量任务与可视化；跨端安装与升级更稳。",
    ),
    (
        "03",
        "加强 AI 理解",
        "自然语言直达可执行计算意图",
        "提升材料术语与 DFTB 语境理解；自动拆解科研任务为计算步骤；课堂中心统一模型，登录即用。",
    ),
]

for i, (num, title, sub, body) in enumerate(cards):
    left = Inches(0.75) + Inches(4.1) * i
    top = Inches(1.85)
    add_rect(s, left, top, Inches(3.9), Inches(4.5), BG)
    add_rect(s, left, top, Inches(3.9), Inches(0.1), ACCENT)
    add_textbox(
        s,
        left + Inches(0.28),
        top + Inches(0.35),
        Inches(3.3),
        Inches(0.4),
        [{"text": num, "size": 18, "bold": True, "color": ACCENT}],
    )
    add_textbox(
        s,
        left + Inches(0.28),
        top + Inches(0.85),
        Inches(3.3),
        Inches(0.55),
        [{"text": title, "size": 22, "bold": True, "color": INK}],
    )
    add_textbox(
        s,
        left + Inches(0.28),
        top + Inches(1.45),
        Inches(3.3),
        Inches(0.7),
        [{"text": sub, "size": 14, "color": ACCENT}],
    )
    add_textbox(
        s,
        left + Inches(0.28),
        top + Inches(2.3),
        Inches(3.3),
        Inches(1.9),
        [{"text": body, "size": 14, "color": INK}],
    )

# footer chrome
add_rect(s, 0, H - Inches(0.42), W, Inches(0.42), BG)
add_rect(s, 0, H - Inches(0.42), W, Pt(1), LINE)
add_textbox(
    s,
    Inches(0.7),
    H - Inches(0.4),
    Inches(9.5),
    Inches(0.36),
    [
        {
            "text": "DFTB-NEU  ·  东北大学材料科学与工程学院 · 张林教授课题组",
            "size": 11,
            "color": MUTED,
        }
    ],
    valign=MSO_ANCHOR.MIDDLE,
)
add_textbox(
    s,
    W - Inches(1.6),
    H - Inches(0.4),
    Inches(1.2),
    Inches(0.36),
    [{"text": "展望", "size": 11, "color": MUTED, "align": PP_ALIGN.RIGHT}],
    valign=MSO_ANCHOR.MIDDLE,
)

desktop = Path.home() / "Desktop"
out = desktop / "DFTB-NEU_总结与展望.pptx"
prs.save(str(out))
print("wrote", out)
